import customtkinter as ctk
from database import Database
from tkinter import messagebox, filedialog, ttk
from tkcalendar import Calendar
import tkinter as tk
import os
import re
from datetime import datetime, timedelta
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import calendar


class DatePickerButton(ctk.CTkFrame):
    def __init__(self, master, initial_date=None):
        super().__init__(master, fg_color="transparent")
        self._date = initial_date or datetime.today().date()
        self._open = False
        self._callback = None

        self.apply_calendar_style()

        self.btn = ctk.CTkButton(
            self,
            text=self._fmt(),
            width=170,
            height=36,
            corner_radius=10,
            fg_color=("#EAECEE", "#1E2A3A"),
            hover_color=("#D5D8DC", "#2C3E50"),
            border_width=1,
            border_color=("#ABB2B9", "#3D5166"),
            text_color=("#1A1A1A", "#FFFFFF"),
            anchor="w",
            command=self.toggle,
        )
        self.btn.pack()

        self.panel = ctk.CTkFrame(
            self.winfo_toplevel(),
            fg_color=("#FFFFFF", "#141E2B"),
            corner_radius=12,
            border_width=1,
            border_color=("#ABB2B9", "#2A3A4A"),
        )

        self.cal = Calendar(
            self.panel,
            style="Custom.Calendar",
            selectmode="day",
            date_pattern="yyyy-mm-dd",
            year=self._date.year,
            month=self._date.month,
            day=self._date.day,
            showweeknumbers=False,
            firstweekday="monday",
            font=("Arial", 11),
            cursor="hand2",
        )
        self.cal.pack(padx=8, pady=8)
        self.cal.bind("<<CalendarSelected>>", self._select)

        self.bind("<Expose>", lambda e: self.apply_calendar_style())

    def apply_calendar_style(self):
        style = ttk.Style()
        style.theme_use("clam")
        is_dark = ctk.get_appearance_mode() == "Dark"

        bg = "#1A1A2E" if is_dark else "#FFFFFF"
        fg = "white" if is_dark else "black"
        h_bg = "#16213E" if is_dark else "#EAECEE"
        h_fg = "#4FC3F7" if is_dark else "#2471A3"

        style.configure(
            "Custom.Calendar",
            background=bg,
            foreground=fg,
            headersbackground=h_bg,
            headersforeground=h_fg,
            selectbackground="#3498DB",
            selectforeground="white",
            normalbackground=bg,
            normalforeground=fg,
            weekendbackground=bg,
            weekendforeground="#F39C12",
            bordercolor=h_bg,
            relief="flat",
        )

        if hasattr(self, "cal"):
            try:
                self.cal.configure(style="Custom.Calendar")
            except Exception:
                pass

    def _fmt(self):
        return f"  📅  {self._date.strftime('%Y-%m-%d')}"

    def toggle(self):
        if self._open:
            self.panel.place_forget()
        else:
            self.panel.lift()
            self.panel.place(in_=self, x=0, y=self.btn.winfo_height() + 2)
        self._open = not self._open

    def _select(self, event):
        selected = self.cal.get_date()
        self._date = datetime.strptime(selected, "%Y-%m-%d").date()
        self.btn.configure(text=self._fmt())
        self.toggle()
        if self._callback:
            self._callback(self._date)

    def get_date(self):
        return self._date

    def set_date(self, d):
        self._date = d
        self.cal.selection_set(d)
        self.btn.configure(text=self._fmt())

    def on_change(self, callback):
        self._callback = callback


class MemberSchedule(ctk.CTkFrame):
    def __init__(self, master, user):
        super().__init__(master, fg_color=("#FFFFFF", "#1A1A1A"))
        self.db = Database()
        self.user = user
        self.view_mode = tk.StringVar(value="List")
        self.overview_year = None
        self.overview_month = None
        self._empty_state_active = False
        self._auto_refresh_job = None
        self.setup_ui()
        self.auto_refresh()

    def setup_ui(self):
        self.wrapper = ctk.CTkFrame(self, fg_color="transparent")
        self.wrapper.pack(fill="both", expand=True, padx=80, pady=5)

        title_row = ctk.CTkFrame(self.wrapper, fg_color="transparent")
        title_row.pack(fill="x", padx=10, pady=(14, 4))
        ctk.CTkLabel(
            title_row,
            text="My WFH / Office Schedule",
            font=("Arial", 22, "bold"),
            text_color=("#1A1A1A", "#FFFFFF"),
        ).pack(side="left")

        filter_card = ctk.CTkFrame(
            self.wrapper,
            corner_radius=14,
            fg_color=("#F2F4F4", "#141E2B"),
            border_width=1,
            border_color=("#D5D8DC", "#253545"),
            height=92,
        )
        filter_card.pack(fill="x", padx=10, pady=(4, 10))
        filter_card.pack_propagate(False)

        inner = ctk.CTkFrame(filter_card, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=18, pady=10)

        def _lbl(parent, text):
            ctk.CTkLabel(
                parent,
                text=text,
                font=("Arial", 11),
                text_color=("#566573", "#8899AA"),
            ).pack(side="left", padx=(0, 4))

        _lbl(inner, "From")
        self.start_cal = DatePickerButton(inner, initial_date=datetime.today().date())
        self.start_cal.pack(side="left", padx=(0, 16))

        _lbl(inner, "To")
        self.end_cal = DatePickerButton(inner, initial_date=datetime.today().date())
        self.end_cal.pack(side="left", padx=(0, 16))

        _lbl(inner, "Status")
        self.status_filter = ctk.CTkComboBox(
            inner,
            values=["All", "Office", "WFH"],
            width=110,
            height=36,
            corner_radius=8,
            border_color=("#ABB2B9", "#3D5166"),
            border_width=1,
            fg_color=("#FFFFFF", "#1E2A3A"),
            button_color=("#D5D8DC", "#2C3E50"),
            text_color=("#1A1A1A", "#FFFFFF"),
            font=("Arial", 12),
        )
        self.status_filter.set("All")
        self.status_filter.pack(side="left", padx=(0, 24))

        spacer = ctk.CTkFrame(inner, fg_color="transparent")
        spacer.pack(side="left", fill="x", expand=True)

        def _btn(parent, text, color, hover, cmd, width=60):
            button = ctk.CTkButton(
                parent,
                text=text,
                width=width,
                height=36,
                corner_radius=8,
                font=("Arial", 12, "bold"),
                fg_color=(color, color),
                hover_color=(hover, hover),
                text_color="#FFFFFF",
                command=cmd,
            )
            button.pack(side="left", padx=3)
            return button

        _btn(inner, "🔍  Filter", "#2471A3", "#1A5276", self.refresh_view)
        _btn(inner, "✖  Clear", "#566573", "#424949", self.clear_filters)
        _btn(inner, "📄  PDF", "#C0392B", "#922B21", self.export_to_pdf)
        _btn(inner, "📥  Excel", "#16A085", "#117A65", self.export_to_xlsx)
        

        view_toggle = ctk.CTkFrame(self.wrapper, fg_color="transparent")
        view_toggle.pack(fill="x", padx=10, pady=(0, 4))
        seg_view = ctk.CTkSegmentedButton(
            view_toggle,
            values=["List", "Overview"],
            variable=self.view_mode,
            font=("Arial", 11, "bold"),
            selected_color=("#2471A3", "#2471A3"),
            unselected_color=("#F0F3F5", "#2C3E50"),
            text_color=("#2C3E50", "#FFFFFF"),
            command=lambda v=None: self.switch_view(v),
        )
        seg_view.pack(side="right")
        try:
            self.view_mode.trace_add("write", lambda *a: self.switch_view())
        except Exception:
            self.view_mode.trace("w", lambda *a: self.switch_view())

        self.overview_frame = ctk.CTkScrollableFrame(
            self.wrapper,
            fg_color=("#FFFFFF", "#0D1117"),
            corner_radius=14,
            border_width=1,
            border_color=("#D5D8DC", "#1E2A3A"),
        )

        self.scroll = ctk.CTkScrollableFrame(
            self.wrapper,
            fg_color=("#FFFFFF", "#0D1117"),
            corner_radius=14,
            border_width=1,
            border_color=("#D5D8DC", "#1E2A3A"),
        )
        self.scroll.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.refresh_view()

    def _show_message(self, message, message_type="info", duration=3000):
        if message_type == "error":
            bg_color = "#E74C3C"
            text_color = "white"
        elif message_type == "warning":
            bg_color = "#F39C12"
            text_color = "white"
        elif message_type == "success":
            bg_color = "#27AE60"
            text_color = "white"
        else:
            bg_color = "#3498DB"
            text_color = "white"

        message_frame = ctk.CTkFrame(
            self.winfo_toplevel(),
            fg_color=bg_color,
            corner_radius=8,
        )
        message_frame.place(relx=1.0, rely=0, x=-20, y=20, anchor="ne")

        ctk.CTkLabel(
            message_frame,
            text=message,
            text_color=text_color,
            font=("Arial", 12, "bold"),
            wraplength=250,
        ).pack(padx=15, pady=10)

        self.winfo_toplevel().after(duration, message_frame.destroy)

    def _timestamped_export_name(self, base_name, default_ext):
        stem, ext = os.path.splitext(base_name)
        if not ext:
            ext = default_ext
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{stem}_{timestamp}{ext}"

    def _ensure_unique_export_path(self, file_path, default_ext):
        directory, filename = os.path.split(file_path)
        stem, ext = os.path.splitext(filename)
        if not ext:
            ext = default_ext
        if re.search(r'_\d{8}_\d{6}$', stem):
            return file_path
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(directory, f"{stem}_{timestamp}{ext}")

    def clear_filters(self, refresh=True):
        self.start_cal.set_date(datetime.today().date())
        self.end_cal.set_date(datetime.today().date())
        self.status_filter.set("All")
        self._empty_state_active = False
        if refresh and self.view_mode.get() == "List":
            self.refresh_view()

    def apply_filters(self):
        if self.view_mode.get() != "List":
            return
        self.refresh_view(show_toast=True)

    def switch_view(self, _=None):
        mode = self.view_mode.get()
        if mode == "Overview":
            self._cancel_auto_refresh()
            try:
                self.scroll.pack_forget()
            except Exception:
                pass
            if self.overview_year is None or self.overview_month is None:
                d = self.start_cal.get_date()
                self.overview_year = d.year
                self.overview_month = d.month
            self.overview_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))
            self.render_overview()
        else:
            try:
                self.overview_frame.pack_forget()
            except Exception:
                pass
            self.scroll.pack(fill="both", expand=True, padx=10, pady=(0, 10))
            self.auto_refresh()

    def prev_month(self):
        if self.overview_year is None or self.overview_month is None:
            return
        month = self.overview_month - 1
        year = self.overview_year
        if month < 1:
            month = 12
            year -= 1
        self.overview_year = year
        self.overview_month = month
        self.render_overview()

    def next_month(self):
        if self.overview_year is None or self.overview_month is None:
            return
        month = self.overview_month + 1
        year = self.overview_year
        if month > 12:
            month = 1
            year += 1
        self.overview_year = year
        self.overview_month = month
        self.render_overview()

    def refresh_view(self, show_toast=True, refresh_overview=False):
        if self.view_mode.get() == "Overview":
            if refresh_overview:
                self.render_overview()
            return

        user_id = self.user.get("id")
        start = self.start_cal.get_date()
        end = self.end_cal.get_date()
        status_search = self.status_filter.get()
        today = datetime.today().date()

        if end < start:
            self._show_message("End date cannot be earlier than Start date", "error")
            return

        try:
            query = (
                "SELECT schedule_date, status FROM wfh_schedules "
                "WHERE user_id=%s AND schedule_date BETWEEN %s AND %s"
            )
            params = [user_id, start, end]
            if status_search != "All":
                query += " AND status=%s"
                params.append(status_search)
            query += " ORDER BY schedule_date ASC"

            self.db.cursor.execute(query, tuple(params))
            rows = self.db.cursor.fetchall()

            if not rows:
                if not show_toast and self._empty_state_active:
                    return
                self._render_empty_state()
                self._empty_state_active = True
                if show_toast:
                    self._show_message("No record found for selected period", "info")
                return

            self._empty_state_active = False
            for child in self.scroll.winfo_children():
                child.destroy()

            header_row = ctk.CTkFrame(
                self.scroll,
                fg_color=("#EAECEE", "#1E2731"),
                corner_radius=12,
            )
            header_row.pack(fill="x", padx=15, pady=(8, 8))
            ctk.CTkLabel(
                header_row,
                text="Date",
                font=("Arial", 12, "bold"),
                text_color=("#2C3E50", "#AABBCD"),
            ).pack(side="left", padx=20, pady=12)
            ctk.CTkLabel(
                header_row,
                text="Status",
                font=("Arial", 12, "bold"),
                text_color=("#2C3E50", "#AABBCD"),
            ).pack(side="right", padx=20, pady=12)

            for r in rows:
                is_today = r["schedule_date"] == today
                is_office = r["status"] == "Office"
                status_color = ("#D5F5E3", "#27AE60") if is_office else ("#D6EAF8", "#3498DB")
                status_text = ("#145A32", "#FFFFFF") if is_office else ("#1B4F72", "#FFFFFF")
                status_icon = "🏢" if is_office else "🏠"

                row_frame = ctk.CTkFrame(
                    self.scroll,
                    fg_color=("#F8F9F9", "#1B2430"),
                    corner_radius=14,
                    border_width=1,
                    border_color=("#EBEDEF", "#26313F"),
                )
                row_frame.pack(fill="x", pady=6, padx=15)

                left = ctk.CTkFrame(row_frame, fg_color="transparent")
                left.pack(side="left", fill="x", expand=True, padx=18, pady=14)

                date_text = r["schedule_date"].strftime("%Y-%m-%d")
                if is_today:
                    date_text += "  ← Today"

                ctk.CTkLabel(
                    left,
                    text=date_text,
                    font=("Arial", 13, "bold"),
                    text_color=("#1A1A1A", "#FFFFFF"),
                    anchor="w",
                ).pack(fill="x")
                ctk.CTkLabel(
                    left,
                    text=r["schedule_date"].strftime("%A"),
                    font=("Arial", 11),
                    text_color=("#566573", "#AABBCD"),
                    anchor="w",
                ).pack(fill="x", pady=(4, 0))

                ctk.CTkLabel(
                    row_frame,
                    text=f"{status_icon} {r['status'].upper()}",
                    fg_color=status_color,
                    text_color=status_text,
                    corner_radius=6,
                    font=("Arial", 11, "bold"),
                    height=30,
                    width=75,
                ).pack(side="right", padx=18, pady=16)
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _render_empty_state(self):
        for child in self.scroll.winfo_children():
            child.destroy()

        ctk.CTkLabel(
            self.scroll,
            text="No schedule assigned for this period.",
            font=("Arial", 13, "bold"),
            text_color=("#667085", "#AABBCD"),
        ).pack(pady=20)

    def render_overview(self):
        for child in self.overview_frame.winfo_children():
            child.destroy()

        user_id = self.user.get("id")
        year = self.overview_year or self.start_cal.get_date().year
        month = self.overview_month or self.start_cal.get_date().month
        self.overview_year = year
        self.overview_month = month

        first_day = datetime(year, month, 1).date()
        last_day_num = calendar.monthrange(year, month)[1]
        last_day = datetime(year, month, last_day_num).date()

        try:
            self.db.cursor.execute(
                """SELECT schedule_date, status
                   FROM wfh_schedules
                   WHERE user_id=%s AND schedule_date BETWEEN %s AND %s
                   ORDER BY schedule_date ASC""",
                (user_id, first_day, last_day),
            )
            rows = self.db.cursor.fetchall()
            summary = {}
            for r in rows:
                summary[r["schedule_date"]] = r["status"]
        except Exception as e:
            messagebox.showerror("Overview Error", str(e))
            summary = {}

        is_dark = ctk.get_appearance_mode() == "Dark"
        month_label_color = ("#2C3E50", "#E6EEF6")
        weekday_text_color = ("#2C3E50", "#BFC9D3")
        date_text_color = ("#0F172A", "#FFFFFF")
        nav_btn_bg = ("#F0F3F5", "#2F3B43")
        nav_btn_text = ("#2C3E50", "#E6EEF6")
        overview_card_bg = ("#F2F4F4", "#0B0F12")
        day_cell_bg = ("#FFFFFF", "#0E1518")
        day_cell_border = ("#CBD5E1", "#24313A")
        day_cell_hover_bg = ("#EAF2F8", "#0F2228")
        day_cell_hover_border = ("#7FB3D5", "#57A6FF")

        nav = ctk.CTkFrame(self.overview_frame, fg_color="transparent")
        nav.pack(fill="x", pady=(6, 8))
        ctk.CTkButton(
            nav,
            text="◀",
            width=36,
            height=30,
            fg_color=nav_btn_bg,
            hover_color=("#E0EAF4", "#3E4A52"),
            text_color=nav_btn_text,
            command=self.prev_month,
        ).pack(side="left")
        ctk.CTkLabel(
            nav,
            text=f"{calendar.month_name[month]} {year}",
            font=("Arial", 18, "bold"),
            text_color=month_label_color,
        ).pack(side="left", padx=12)
        ctk.CTkButton(
            nav,
            text="▶",
            width=36,
            height=30,
            fg_color=nav_btn_bg,
            hover_color=("#E0EAF4", "#3E4A52"),
            text_color=nav_btn_text,
            command=self.next_month,
        ).pack(side="left")

        calendar_wrap = ctk.CTkFrame(
            self.overview_frame,
            fg_color=overview_card_bg,
            corner_radius=14,
            border_width=1,
            border_color=("#D5D8DC", "#1E2A3A"),
        )
        calendar_wrap.pack(fill="both", expand=True, pady=12, padx=(12, 24))

        weekdays = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        day_weights = [5, 5, 5, 5, 5, 0, 0]
        day_minsizes = [0, 0, 0, 0, 0, 72, 72]
        day_pads = [6, 6, 6, 6, 6, 1, 1]
        weekend_width = 72

        for idx in range(7):
            calendar_wrap.grid_columnconfigure(idx, weight=day_weights[idx], minsize=day_minsizes[idx])

        for idx, wd in enumerate(weekdays):
            label_width = weekend_width if idx >= 5 else 1
            ctk.CTkLabel(
                calendar_wrap,
                text=wd,
                anchor="center",
                font=("Arial", 13, "bold"),
                text_color=weekday_text_color,
                width=label_width,
            ).grid(row=0, column=idx, sticky="ew", padx=(day_pads[idx], day_pads[idx]), pady=(0, 6))

        month_matrix = calendar.monthcalendar(year, month)

        for r_idx, week in enumerate(month_matrix, start=1):
            calendar_wrap.grid_rowconfigure(r_idx, weight=1)
            for c_idx, day in enumerate(week):
                cell = ctk.CTkFrame(
                    calendar_wrap,
                    width=weekend_width if c_idx >= 5 else 1,
                    height=120,
                    fg_color=day_cell_bg,
                    corner_radius=10,
                    border_width=1,
                    border_color=day_cell_border,
                )
                cell.grid(
                    row=r_idx,
                    column=c_idx,
                    sticky="nsew",
                    padx=(day_pads[c_idx], day_pads[c_idx]),
                    pady=4,
                )
                cell.pack_propagate(False)

                if day == 0:
                    placeholder = ctk.CTkFrame(cell, fg_color=day_cell_bg)
                    placeholder.pack(fill="both", expand=True)
                    continue

                d_date = datetime(year, month, day).date()
                is_today = d_date == datetime.today().date()
                status = summary.get(d_date)

                def _on_enter(ev, widget=cell):
                    try:
                        widget.configure(border_color=day_cell_hover_border)
                        widget.configure(fg_color=day_cell_hover_bg)
                    except Exception:
                        pass

                def _on_leave(ev, widget=cell):
                    try:
                        widget.configure(border_color=day_cell_border)
                        widget.configure(fg_color=day_cell_bg)
                    except Exception:
                        pass

                cell.bind("<Enter>", _on_enter)
                cell.bind("<Leave>", _on_leave)

                top_row = ctk.CTkFrame(cell, fg_color="transparent")
                top_row.pack(fill="x", padx=10, pady=(8, 2))
                date_txt = str(day) + ("  ← Today" if is_today else "")
                ctk.CTkLabel(
                    top_row,
                    text=date_txt,
                    anchor="w",
                    font=("Arial", 14, "bold"),
                    text_color=date_text_color,
                ).pack(side="left")

                counts_frame = ctk.CTkFrame(cell, fg_color="transparent")
                counts_frame.pack(fill="both", expand=True, padx=10, pady=6)

                if status:
                    pill_bg = ("#0F3B4A", "#0F3B4A") if status == "WFH" else ("#0F3A24", "#0F3A24")
                    pill_text = ("#8FD8FF", "#8FD8FF") if status == "WFH" else ("#9EE6B6", "#9EE6B6")
                    icon = "🏠" if status == "WFH" else "🏢"
                    ctk.CTkLabel(
                        counts_frame,
                        text=f"  {icon} {status}  ",
                        font=("Arial", 12),
                        fg_color=pill_bg,
                        text_color=pill_text,
                        corner_radius=8,
                    ).pack(anchor="w", pady=4)
                elif is_today:
                    ctk.CTkLabel(
                        counts_frame,
                        text="No schedule yet",
                        font=("Arial", 11),
                        text_color=("#7B7D7D", "#AABBCD"),
                    ).pack(anchor="w", pady=4)

    def export_to_pdf(self):
        user_id = self.user.get("id")
        start = self.start_cal.get_date()
        end = self.end_cal.get_date()

        file_path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            initialfile=self._timestamped_export_name("My_Schedule.pdf", ".pdf"),
        )
        if not file_path:
            return

        try:
            selected_path = file_path
            file_path = self._ensure_unique_export_path(selected_path, ".pdf")

            self.db.cursor.execute(
                """SELECT schedule_date, status
                   FROM wfh_schedules
                   WHERE user_id=%s AND schedule_date BETWEEN %s AND %s
                   ORDER BY schedule_date ASC""",
                (user_id, start, end),
            )
            data = self.db.cursor.fetchall()

            doc = SimpleDocTemplate(file_path, pagesize=letter)
            styles = getSampleStyleSheet()
            line_style = ParagraphStyle(
                "MemberLine",
                parent=styles["Normal"],
                leftIndent=0,
            )
            elements = [
                Paragraph(f"My Work Schedule", styles["Title"]),
                Spacer(1, 8),
                Paragraph(f"Date: {start} ~ {end}", line_style),
                Spacer(1, 15),
            ]
            table_data = [["Date", "Member Name", "Status"]]
            member_name = self.user.get("full_name") or self.user.get("username") or self.user.get("name") or "Member"
            for r in data:
                day = r["schedule_date"].strftime("%a")
                table_data.append([f"{r['schedule_date']} ({day})", member_name, r["status"]])

            t = Table(table_data, colWidths=[130, 200, 100])
            t.hAlign = "LEFT"
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ]
                )
            )
            elements.append(t)
            doc.build(elements)
            self._show_message("PDF exported successfully!", "success", duration=3000)
        except Exception as e:
            self._show_message(f"PDF Error: {e}", "error", duration=3000)

    def export_to_xlsx(self):
        user_id = self.user.get("id")
        member_name = self.user.get("full_name") or self.user.get("username") or self.user.get("name") or "Member"
        start = self.start_cal.get_date()
        end = self.end_cal.get_date()

        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=self._timestamped_export_name("My_Schedule.xlsx", ".xlsx"),
        )
        if not file_path:
            return

        try:
            selected_path = file_path
            file_path = self._ensure_unique_export_path(selected_path, ".xlsx")

            self.db.cursor.execute(
                """SELECT schedule_date, status
                   FROM wfh_schedules
                   WHERE user_id=%s AND schedule_date BETWEEN %s AND %s
                   ORDER BY schedule_date ASC""",
                (user_id, start, end),
            )
            data = self.db.cursor.fetchall()

            import openpyxl
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "Schedule"

            ws.append(["My Work Schedule"])
            ws.append([f"Date: {start} ~ {end}"])
            ws.append([])
            headers = ["Date", "Member Name", "Status"]
            ws.append(headers)
            for r in data:
                dstr = r["schedule_date"].strftime("%Y-%m-%d") if hasattr(r["schedule_date"], "strftime") else str(r["schedule_date"])
                ws.append([dstr, member_name, r["status"]])

            thin = Side(border_style="thin", color="444444")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            col_widths = [0] * ws.max_column
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False), start=1):
                for col_idx, cell in enumerate(row, start=1):
                    if row_idx >= 4:
                        cell.border = border
                    if row_idx == 4:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(fill_type="solid", fgColor="DDDDDD")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif row_idx == 1:
                        cell.font = Font(bold=True, size=16)
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif row_idx in (2, 3):
                        cell.font = Font(bold=False, size=11)
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    val = str(cell.value) if cell.value is not None else ""
                    col_widths[col_idx - 1] = max(col_widths[col_idx - 1], len(val))

            ws.merge_cells("A1:C1")
            ws.merge_cells("A2:C2")

            for i, wth in enumerate(col_widths, start=1):
                col = get_column_letter(i)
                ws.column_dimensions[col].width = max(10, wth + 2)

            wb.save(file_path)
            self._show_message("Excel exported successfully!", "success", duration=3000)
        except Exception as e:
            self._show_message(f"Excel Error: {e}", "error", duration=3000)

    def auto_refresh(self):
        if self.winfo_exists():
            if self.view_mode.get() != "List":
                self._cancel_auto_refresh()
                return
            self.refresh_view(show_toast=False)
            self._auto_refresh_job = self.after(30000, self.auto_refresh)

    def _cancel_auto_refresh(self):
        if self._auto_refresh_job is not None:
            try:
                self.after_cancel(self._auto_refresh_job)
            except Exception:
                pass
            self._auto_refresh_job = None
